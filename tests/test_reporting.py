from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from fpl_dashboard.processing import DEMAND_COLUMNS, ENERGY_COLUMNS
from fpl_dashboard.reporting import create_excel_report


def _monthly_summary():
    rows = []
    for account in ["A", "B"]:
        for month, source in [(1, "Actual"), (2, "Estimated")]:
            row = {
                "Account": account,
                "Year": 2025,
                "Month": month,
                "Data Source": source,
                "Estimate Method": "Actual uploaded interval file" if source == "Actual" else "Interpolated",
                "Confidence": "Normal",
                "Peak During Non-Operating": False,
                "Non-Operating %": 20.0,
                "Coverage %": 100.0,
                "Coverage Status": "Complete",
                "Uploaded Row Count": 2880,
                "Expected Row Count": 2880,
                "Total Rows": 2880,
                "Operating Rows": 2400,
                "Not Operating Rows": 480,
            }
            for column in ENERGY_COLUMNS + DEMAND_COLUMNS:
                row[column] = 100.0
            rows.append(row)
    return pd.DataFrame(rows)


def test_excel_report_contains_required_sheets_with_estimates():
    file_log = pd.DataFrame({"Account": ["A"], "File name": ["jan.xlsx"], "Status": ["Valid"]})
    notes = pd.DataFrame({"Account number": ["A"], "Estimated month": ["February"]})
    content = create_excel_report(_monthly_summary(), file_log, notes)
    workbook = pd.ExcelFile(BytesIO(content))
    required = {
        "Monthly Summary",
        "Official Dashboard",
        "Data Quality",
        "Operating vs Non-Operating",
        "On-Peak vs Off-Peak Summary",
        "Demand Summary",
        "Account-Level Summary",
        "Consolidated Summary",
        "Classification Audit",
        "Daily Hourly Breakdown",
        "Input File Log",
        "Estimation Notes",
        "Chart Data",
    }
    assert required.issubset(set(workbook.sheet_names))
    assert workbook.sheet_names[:2] == ["Official Dashboard", "Monthly Summary"]

    styled_workbook = load_workbook(BytesIO(content))
    official_sheet = styled_workbook["Official Dashboard"]
    assert official_sheet["A1"].fill.fgColor.rgb == "FF95B3D7"
    assert official_sheet.max_row == len(_monthly_summary()) + 2
    total_row = [cell.value for cell in official_sheet[official_sheet.max_row]]
    assert "Total" in total_row
    total_kwh_column = [cell.value for cell in official_sheet[1]].index("Total kWh") + 1
    assert official_sheet.cell(official_sheet.max_row, total_kwh_column).fill.fgColor.rgb == "FFFFFF00"

    monthly = pd.read_excel(BytesIO(content), sheet_name="Monthly Summary")
    assert "Estimated" in set(monthly["Data Source"])
    official = pd.read_excel(BytesIO(content), sheet_name="Official Dashboard")
    assert "On Peak Operating (kWh)" in official.columns
    quality = pd.read_excel(BytesIO(content), sheet_name="Data Quality")
    assert "Coverage Status" in quality.columns
    audit = pd.read_excel(BytesIO(content), sheet_name="Classification Audit")
    assert "Operating Rows" in audit.columns
